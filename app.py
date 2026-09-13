# -*- coding: utf-8 -*-

import io
import re
import unicodedata
from pathlib import Path

import pandas as pd
import requests
import streamlit as st

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter


st.set_page_config(
    page_title="農薬散布記録一覧",
    layout="wide",
)


PESTICIDE_GROUP_ORDER = ["殺菌剤", "殺虫剤", "展着剤", "除草剤"]
PESTICIDE_GROUP_RANK = {name: i for i, name in enumerate(PESTICIDE_GROUP_ORDER)}

REGISTRATION_COLUMNS = [
    "希釈倍数使用量",
    "使用時期",
    "本剤の使用回数",
]

OPTIONAL_DISPLAY_COLUMNS = [
    {"key": "農薬グループ", "label": "農薬グループ"},
    {"key": "農薬名", "label": "農薬名"},
    {"key": "有効成分", "label": "有効成分"},
    {"key": "作用機作分類", "label": "作用機作分類"},
    {"key": "希釈倍数使用量", "label": "希釈倍率"},
    {"key": "使用時期", "label": "使用時期"},
    {"key": "本剤の使用回数", "label": "使用回数"},
]


def read_csv_auto(file_or_path) -> pd.DataFrame:
    """アップロードCSVまたはローカルCSVを文字コード自動判定で読む"""
    encodings = ["utf-8-sig", "cp932", "shift_jis", "utf-8"]

    if hasattr(file_or_path, "read"):
        if hasattr(file_or_path, "seek"):
            file_or_path.seek(0)
        raw = file_or_path.read()
    else:
        with open(file_or_path, "rb") as f:
            raw = f.read()

    last_error = None

    for enc in encodings:
        try:
            return pd.read_csv(
                io.BytesIO(raw),
                encoding=enc,
                dtype=str,
                keep_default_na=False,
            )
        except Exception as e:
            last_error = e

    raise RuntimeError(f"CSVを読み込めませんでした: {last_error}")


@st.cache_data(ttl=60, show_spinner=False)
def download_google_drive_csv(file_id: str) -> bytes:
    """リンク共有されたGoogle Drive上のCSVを取得する。"""
    file_id = str(file_id).strip()
    if not file_id:
        raise ValueError("Google DriveのファイルIDが空です。")

    response = requests.get(
        "https://drive.google.com/uc",
        params={"export": "download", "id": file_id},
        timeout=30,
    )
    response.raise_for_status()

    content_type = response.headers.get("Content-Type", "").lower()
    content_start = response.content.lstrip()[:100].lower()
    if "text/html" in content_type or content_start.startswith(b"<!doctype html"):
        raise RuntimeError(
            "Google DriveからCSVではなくWebページが返されました。"
            "共有設定を「リンクを知っている全員・閲覧者」にしてください。"
        )

    return response.content


def read_google_drive_csv(file_id: str) -> pd.DataFrame:
    """Google Drive上のCSVを文字コード自動判定で読む。"""
    return read_csv_auto(io.BytesIO(download_google_drive_csv(file_id)))


def get_drive_file_ids():
    """Streamlit Secretsから3種類のCSVのファイルIDを取得する。"""
    try:
        drive = st.secrets.get("google_drive", {})
        return {
            "main": str(drive.get("main_csv_file_id", "")).strip(),
            "pesticide": str(drive.get("pesticide_info_file_id", "")).strip(),
            "registration": str(drive.get("registration_info_file_id", "")).strip(),
        }
    except Exception:
        return {"main": "", "pesticide": "", "registration": ""}


def clean_text(value):
    """None、nan、NaT、null などを空白にする"""
    if pd.isna(value):
        return ""

    value = str(value).strip()

    if value.lower() in ["none", "nan", "nat", "null"]:
        return ""

    return value


def normalize_id(value):
    """農薬IDを結合用に整える"""
    value = clean_text(value)
    value = unicodedata.normalize("NFKC", value)

    if value.endswith(".0") and value[:-2].isdigit():
        value = value[:-2]

    return value


def normalize_registration_no(value):
    """登録番号を結合用に整える。例：第24185号 → 24185"""
    value = clean_text(value)
    value = unicodedata.normalize("NFKC", value)
    value = value.replace("第", "").replace("号", "")
    value = value.replace(" ", "").replace("　", "")

    if value.endswith(".0") and value[:-2].isdigit():
        value = value[:-2]

    digits = re.findall(r"\d+", value)
    if digits:
        return "".join(digits)

    return value


def normalize_pesticide_group(value):
    """農薬グループを整える。空白は未分類にする"""
    value = clean_text(value)
    return value if value else "未分類"


def pesticide_group_sort_key(value):
    value = normalize_pesticide_group(value)
    return (PESTICIDE_GROUP_RANK.get(value, 99), value)


def sorted_pesticide_groups(groups):
    return sorted(
        {normalize_pesticide_group(g) for g in groups if normalize_pesticide_group(g)},
        key=pesticide_group_sort_key,
    )


def join_unique(values):
    """重複を除いて / 区切りで連結する"""
    result = []
    seen = set()

    for value in values:
        text = clean_text(value)
        if text and text not in seen:
            result.append(text)
            seen.add(text)

    return " / ".join(result)


def format_date(dt):
    """日付を m/d 形式にする"""
    if pd.isna(dt):
        return ""

    dt = pd.Timestamp(dt)
    return f"{dt.month}/{dt.day}"


def extract_allowed_use_count(value):
    """
    本剤の使用回数から最大回数を取り出す。
    例：
    4回以内 → 4
    3回以内 / 4回以内 → 4
    """
    text = clean_text(value)
    if not text:
        return None

    text = unicodedata.normalize("NFKC", text)

    numbers = []

    for match in re.findall(r"(\d+)\s*回\s*以内", text):
        try:
            numbers.append(int(match))
        except ValueError:
            pass

    if not numbers:
        for match in re.findall(r"(\d+)\s*回", text):
            try:
                numbers.append(int(match))
            except ValueError:
                pass

    if not numbers:
        return None

    return max(numbers)


def prepare_pesticide_info(info_df: pd.DataFrame) -> pd.DataFrame:
    """農薬情報CSVを結合しやすい形に整える"""

    required_cols = ["農薬ID", "農薬グループ", "有効成分", "登録番号"]
    missing = [c for c in required_cols if c not in info_df.columns]

    if missing:
        raise ValueError(f"農薬情報CSVに必要な列がありません: {missing}")

    info = info_df.copy()

    info["農薬ID_key"] = info["農薬ID"].apply(normalize_id)
    info["登録番号_key"] = info["登録番号"].apply(normalize_registration_no)
    info["農薬グループ"] = info["農薬グループ"].apply(normalize_pesticide_group)
    info["有効成分"] = info["有効成分"].apply(clean_text)

    # 農薬情報.csv の「メモ」を、表では「作用機作分類」として扱う
    if "メモ" in info.columns:
        info["作用機作分類"] = info["メモ"].apply(clean_text)
    elif "作用機作分類" in info.columns:
        info["作用機作分類"] = info["作用機作分類"].apply(clean_text)
    else:
        info["作用機作分類"] = ""

    if "農薬名" in info.columns:
        info["農薬名_情報"] = info["農薬名"].apply(clean_text)
    else:
        info["農薬名_情報"] = ""

    info = info[info["農薬ID_key"] != ""].copy()
    info = info.drop_duplicates(subset=["農薬ID_key"], keep="first")

    return info[
        [
            "農薬ID_key",
            "登録番号_key",
            "農薬名_情報",
            "農薬グループ",
            "有効成分",
            "作用機作分類",
        ]
    ]


def prepare_registration_info(reg_df: pd.DataFrame) -> pd.DataFrame:
    """農薬登録情報CSVを登録番号単位に集約する"""

    required_cols = ["登録番号"] + REGISTRATION_COLUMNS
    missing = [c for c in required_cols if c not in reg_df.columns]

    if missing:
        raise ValueError(f"農薬登録情報CSVに必要な列がありません: {missing}")

    reg = reg_df.copy()
    reg["登録番号_key"] = reg["登録番号"].apply(normalize_registration_no)

    for col in REGISTRATION_COLUMNS:
        reg[col] = reg[col].apply(clean_text)

    reg = reg[reg["登録番号_key"] != ""].copy()

    aggregated = (
        reg.groupby("登録番号_key", dropna=False)[REGISTRATION_COLUMNS]
        .agg(join_unique)
        .reset_index()
    )

    return aggregated


def prepare_base_dataframe(
    main_df: pd.DataFrame,
    pesticide_info_df: pd.DataFrame,
    registration_info_df: pd.DataFrame,
):
    """メインCSV、農薬情報CSV、農薬登録情報CSVを結合して基本整形する"""

    required_cols = ["日付", "作付名", "圃場グループ", "農薬ID", "農薬名"]
    missing = [c for c in required_cols if c not in main_df.columns]

    if missing:
        raise ValueError(f"メインCSVに必要な列がありません: {missing}")

    df = main_df.copy()

    df["日付"] = pd.to_datetime(df["日付"], errors="coerce")
    df = df.dropna(subset=["日付"])

    if df.empty:
        raise ValueError("日付を読み取れるデータがありません。")

    df["作付名"] = df["作付名"].apply(clean_text)
    df["圃場グループ"] = df["圃場グループ"].apply(clean_text)
    df["農薬名"] = df["農薬名"].apply(clean_text)
    df["農薬ID_key"] = df["農薬ID"].apply(normalize_id)

    df = df[(df["作付名"] != "") & (df["農薬ID_key"] != "")].copy()

    if df.empty:
        raise ValueError("作付名・農薬IDを読み取れるデータがありません。")

    info = prepare_pesticide_info(pesticide_info_df)
    reg = prepare_registration_info(registration_info_df)

    merged = df.merge(info, on="農薬ID_key", how="left")

    unmatched_pesticide_info_count = int(
        merged["農薬グループ"].apply(clean_text).eq("").sum()
    )

    reg_keys = set(reg["登録番号_key"].dropna().astype(str))
    unmatched_registration_count = int(
        (
            merged["登録番号_key"].apply(clean_text).ne("")
            & ~merged["登録番号_key"].astype(str).isin(reg_keys)
        ).sum()
    )

    merged = merged.merge(reg, on="登録番号_key", how="left")

    merged["農薬名"] = merged.apply(
        lambda row: clean_text(row.get("農薬名_情報", ""))
        or clean_text(row.get("農薬名", "")),
        axis=1,
    )

    merged["農薬グループ"] = merged["農薬グループ"].apply(normalize_pesticide_group)
    merged["有効成分"] = merged["有効成分"].apply(clean_text)
    merged["作用機作分類"] = merged["作用機作分類"].apply(clean_text)

    for col in REGISTRATION_COLUMNS:
        merged[col] = merged[col].apply(clean_text)

    merged = merged[(merged["作付名"] != "") & (merged["農薬名"] != "")].copy()

    if merged.empty:
        raise ValueError("集計できる農薬散布データがありません。")

    return merged, unmatched_pesticide_info_count, unmatched_registration_count


def prepare_summary(df: pd.DataFrame, start_date, end_date):
    """指定期間で農薬散布一覧を作成"""

    start_ts = pd.Timestamp(start_date)
    end_ts = pd.Timestamp(end_date)

    filtered = df[
        (df["日付"] >= start_ts)
        & (df["日付"] <= end_ts)
    ].copy()

    if filtered.empty:
        raise ValueError("指定期間内の農薬散布データがありません。")

    base_columns = [
        "圃場グループ",
        "作付名",
        "農薬グループ",
        "農薬名",
        "有効成分",
        "作用機作分類",
        "希釈倍数使用量",
        "使用時期",
        "本剤の使用回数",
        "農薬ID_key",
        "日付",
    ]

    base = (
        filtered[base_columns]
        .drop_duplicates()
        .sort_values(["圃場グループ", "作付名", "農薬名", "日付"])
    )

    rows = []

    group_columns = [
        "圃場グループ",
        "作付名",
        "農薬グループ",
        "農薬名",
        "有効成分",
        "作用機作分類",
        "希釈倍数使用量",
        "使用時期",
        "本剤の使用回数",
        "農薬ID_key",
    ]

    for group_values, g in base.groupby(group_columns, dropna=False):
        (
            field_group,
            crop_name,
            pesticide_group,
            pesticide_name,
            active_ingredient,
            mode_of_action,
            dilution_amount,
            timing,
            use_count,
            pesticide_id,
        ) = group_values

        dates = sorted(g["日付"].dropna().unique())

        row = {
            "圃場グループ": clean_text(field_group),
            "作付名": clean_text(crop_name),
            "農薬グループ": normalize_pesticide_group(pesticide_group),
            "農薬名": clean_text(pesticide_name),
            "有効成分": clean_text(active_ingredient),
            "作用機作分類": clean_text(mode_of_action),
            "希釈倍数使用量": clean_text(dilution_amount),
            "使用時期": clean_text(timing),
            "本剤の使用回数": clean_text(use_count),
            "_使用可能回数": extract_allowed_use_count(use_count),
        }

        for i, d in enumerate(dates, start=1):
            row[f"{i}回目"] = format_date(d)

        rows.append(row)

    summary = pd.DataFrame(rows)

    if summary.empty:
        raise ValueError("集計できる農薬散布データがありません。")

    summary = summary.fillna("").replace(["None", "nan", "NaT", "null"], "")

    count_cols = [c for c in summary.columns if c.endswith("回目")]
    max_count = len(count_cols)

    summary["_農薬グループ順"] = summary["農薬グループ"].apply(
        lambda x: pesticide_group_sort_key(x)[0]
    )
    summary["_農薬グループ名順"] = summary["農薬グループ"].apply(
        lambda x: pesticide_group_sort_key(x)[1]
    )

    summary = summary.sort_values(
        [
            "圃場グループ",
            "作付名",
            "_農薬グループ順",
            "_農薬グループ名順",
            "農薬名",
            "有効成分",
            "作用機作分類",
        ]
    ).reset_index(drop=True)

    summary = summary.drop(columns=["_農薬グループ順", "_農薬グループ名順"])

    return summary, max_count


def clean_for_excel(value):
    value = clean_text(value)
    return "" if value in ["None", "nan", "NaT", "null"] else value


def get_display_column_keys(selected_optional_columns):
    """実際に表示する内部列名"""
    return ["作付名"] + selected_optional_columns


def get_display_column_labels(selected_optional_columns):
    """表示用の列名"""
    label_map = {item["key"]: item["label"] for item in OPTIONAL_DISPLAY_COLUMNS}
    labels = ["作付名"]

    for col in selected_optional_columns:
        labels.append(label_map.get(col, col))

    return labels


def get_column_width(header_name):
    """列名に応じてExcel列幅を返す"""
    if header_name == "作付名":
        return 24
    if header_name == "農薬グループ":
        return 14
    if header_name == "農薬名":
        return 24
    if header_name == "有効成分":
        return 36
    if header_name == "作用機作分類":
        return 28
    if header_name == "希釈倍率":
        return 28
    if header_name == "使用時期":
        return 22
    if header_name == "使用回数":
        return 22
    if header_name.endswith("回目"):
        return 11
    return 16


def apply_sheet_style(
    ws,
    header_row: int,
    freeze_cell: str,
    display_rows=None,
    display_cols=None,
    date_start_col=None,
    max_count=0,
):
    """Excelシートの装飾"""

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    title_fill = PatternFill("solid", fgColor="FFF2CC")
    crop_fill = PatternFill("solid", fgColor="E2F0D9")
    over_limit_fill = PatternFill("solid", fgColor="595959")
    over_limit_font = Font(name="Yu Gothic", size=10, color="FFFFFF")

    max_row = display_rows if display_rows is not None else ws.max_row
    max_col = display_cols if display_cols is not None else ws.max_column

    ws.freeze_panes = freeze_cell

    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)

            if cell.value in [None, "None", "nan", "NaT", "null"]:
                cell.value = ""

            cell.border = border
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            cell.font = Font(name="Yu Gothic", size=10)

    for cell in ws[header_row]:
        cell.fill = header_fill
        cell.font = Font(name="Yu Gothic", size=10, bold=True)

    for row in range(1, header_row):
        for col in range(1, min(max_col, 2) + 1):
            ws.cell(row=row, column=col).fill = title_fill
            ws.cell(row=row, column=col).font = Font(name="Yu Gothic", size=10, bold=True)

    for row in range(header_row + 1, ws.max_row + 1):
        ws.cell(row=row, column=1).fill = crop_fill

    # 使用回数超過部分を暗く塗る
    if date_start_col is not None and max_count > 0:
        helper_col = max_col + 1

        for row in range(header_row + 1, ws.max_row + 1):
            allowed_count = ws.cell(row=row, column=helper_col).value

            try:
                allowed_count = int(allowed_count)
            except Exception:
                allowed_count = None

            if allowed_count is not None:
                for i in range(1, max_count + 1):
                    target_col = date_start_col + i - 1

                    if i > allowed_count:
                        cell = ws.cell(row=row, column=target_col)
                        cell.fill = over_limit_fill
                        cell.font = over_limit_font

    # 補助列を非表示
    helper_col = max_col + 1
    if helper_col <= ws.max_column:
        ws.column_dimensions[get_column_letter(helper_col)].hidden = True

    for col in range(1, max_col + 1):
        col_letter = get_column_letter(col)
        header_name = clean_text(ws.cell(row=header_row, column=col).value)
        ws.column_dimensions[col_letter].width = get_column_width(header_name)

    for row in range(1, max_row + 1):
        ws.row_dimensions[row].height = 28


def merge_same_values(ws, col_num: int, start_row: int):
    """同じ値が連続するセルを縦結合する"""
    max_row = ws.max_row

    if max_row < start_row:
        return

    current_value = ws.cell(start_row, col_num).value
    merge_start = start_row

    for row in range(start_row + 1, max_row + 2):
        value = ws.cell(row, col_num).value if row <= max_row else None

        if value != current_value:
            if merge_start < row - 1 and current_value not in [None, ""]:
                ws.merge_cells(
                    start_row=merge_start,
                    start_column=col_num,
                    end_row=row - 1,
                    end_column=col_num,
                )
                ws.cell(merge_start, col_num).alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )

            merge_start = row
            current_value = value


def write_table(
    ws,
    df: pd.DataFrame,
    max_count: int,
    header_row: int,
    start_row: int,
    selected_optional_columns: list[str],
):
    """Excelに一覧表を書き込む"""

    display_column_keys = get_display_column_keys(selected_optional_columns)
    display_column_labels = get_display_column_labels(selected_optional_columns)

    headers = display_column_labels + [
        f"{i}回目" for i in range(1, max_count + 1)
    ]

    helper_col = len(headers) + 1

    for col, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=h)

    ws.cell(row=header_row, column=helper_col, value="_使用可能回数")

    for r_idx, (_, row) in enumerate(df.iterrows(), start=start_row):
        values = [clean_for_excel(row.get(col, "")) for col in display_column_keys]

        for i in range(1, max_count + 1):
            values.append(clean_for_excel(row.get(f"{i}回目", "")))

        for c_idx, value in enumerate(values, start=1):
            ws.cell(row=r_idx, column=c_idx, value=value)

        allowed_count = row.get("_使用可能回数", None)
        if allowed_count == "":
            allowed_count = None

        ws.cell(row=r_idx, column=helper_col, value=allowed_count)


def create_excel(
    summary: pd.DataFrame,
    max_count: int,
    start_date,
    end_date,
    selected_field_group: str,
    selected_pesticide_groups: list[str],
    selected_optional_columns: list[str],
) -> bytes:
    """Excelファイルを作成して bytes で返す"""

    summary = summary.fillna("").replace(["None", "nan", "NaT", "null"], "")

    filtered_summary = summary[
        summary["農薬グループ"].isin(selected_pesticide_groups)
    ].copy()

    group_df = filtered_summary[
        filtered_summary["圃場グループ"] == selected_field_group
    ].copy()

    wb = Workbook()
    wb.remove(wb.active)

    # =========================
    # グループ別一覧
    # =========================
    group_ws = wb.create_sheet("グループ別一覧")

    group_ws["A1"] = "対象期間"
    group_ws["B1"] = (
        f"{pd.Timestamp(start_date).strftime('%Y/%m/%d')} 〜 "
        f"{pd.Timestamp(end_date).strftime('%Y/%m/%d')}"
    )
    group_ws["A2"] = "圃場グループ"
    group_ws["B2"] = clean_text(selected_field_group)
    group_ws["A3"] = "農薬グループ"
    group_ws["B3"] = "、".join(selected_pesticide_groups)

    write_table(
        group_ws,
        group_df,
        max_count=max_count,
        header_row=5,
        start_row=6,
        selected_optional_columns=selected_optional_columns,
    )

    display_cols = 1 + len(selected_optional_columns) + max_count
    date_start_col = 1 + len(selected_optional_columns) + 1
    freeze_cell = f"{get_column_letter(date_start_col)}6"

    apply_sheet_style(
        group_ws,
        header_row=5,
        freeze_cell=freeze_cell,
        display_rows=max(group_ws.max_row, 80),
        display_cols=display_cols,
        date_start_col=date_start_col,
        max_count=max_count,
    )

    merge_same_values(group_ws, col_num=1, start_row=6)

    # =========================
    # 一覧データ
    # =========================
    data_ws = wb.create_sheet("一覧データ")

    data_headers = ["圃場グループ"] + get_display_column_labels(selected_optional_columns) + [
        f"{i}回目" for i in range(1, max_count + 1)
    ] + ["_使用可能回数"]

    for col_idx, h in enumerate(data_headers, start=1):
        data_ws.cell(row=1, column=col_idx, value=h)

    display_column_keys = get_display_column_keys(selected_optional_columns)

    for r_idx, (_, row) in enumerate(filtered_summary.iterrows(), start=2):
        values = [clean_for_excel(row.get("圃場グループ", ""))]
        values += [clean_for_excel(row.get(col, "")) for col in display_column_keys]

        for i in range(1, max_count + 1):
            values.append(clean_for_excel(row.get(f"{i}回目", "")))

        values.append(row.get("_使用可能回数", ""))

        for c_idx, value in enumerate(values, start=1):
            data_ws.cell(row=r_idx, column=c_idx, value=value)

    data_display_cols = 1 + display_cols
    data_date_start_col = 1 + date_start_col
    data_freeze_cell = f"{get_column_letter(data_date_start_col)}2"

    apply_sheet_style(
        data_ws,
        header_row=1,
        freeze_cell=data_freeze_cell,
        display_rows=max(data_ws.max_row, 80),
        display_cols=data_display_cols,
        date_start_col=data_date_start_col,
        max_count=max_count,
    )

    wb._sheets = [
        wb["グループ別一覧"],
        wb["一覧データ"],
    ]

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return buffer.getvalue()


def rename_for_display(df: pd.DataFrame) -> pd.DataFrame:
    """Web表示用に列名を変更"""
    rename_map = {
        "希釈倍数使用量": "希釈倍率",
        "本剤の使用回数": "使用回数",
    }
    return df.rename(columns=rename_map)


def make_display_df(df: pd.DataFrame) -> pd.DataFrame:
    """Web表示用に None などを空白にする"""
    display_df = df.copy()
    display_df = display_df.fillna("")
    display_df = display_df.replace(["None", "nan", "NaT", "null"], "")

    if "_使用可能回数" in display_df.columns:
        display_df = display_df.drop(columns=["_使用可能回数"])

    display_df = rename_for_display(display_df)

    return display_df


def style_web_table(display_df: pd.DataFrame, original_df: pd.DataFrame, max_count: int):
    """Web表示で使用回数超過部分を暗く塗る"""

    date_cols = [f"{i}回目" for i in range(1, max_count + 1)]

    def style_row(row):
        styles = [""] * len(row)

        allowed_count = None

        if row.name in original_df.index:
            raw_allowed = original_df.loc[row.name, "_使用可能回数"]
            try:
                allowed_count = int(raw_allowed)
            except Exception:
                allowed_count = None

        if allowed_count is not None:
            for i, col in enumerate(date_cols, start=1):
                if col in row.index and i > allowed_count:
                    col_idx = list(row.index).index(col)
                    styles[col_idx] = (
                        "background-color: #595959; "
                        "color: white;"
                    )

        return styles

    return display_df.style.apply(style_row, axis=1)


def display_table_like_excel(df: pd.DataFrame, max_count: int):
    """Web上で作付名を結合っぽく見せ、使用不可回数を暗く表示する"""

    original_df = df.copy()
    display_df = make_display_df(df)
    column_config = {}

    if "作付名" in display_df.columns:
        def display_units(value):
            return sum(
                2 if unicodedata.east_asian_width(char) in {"W", "F", "A"} else 1
                for char in clean_text(value)
            )

        longest_name = max(
            [display_units("作付名")]
            + [display_units(value) for value in display_df["作付名"]]
        )
        crop_name_width = max(200, min(600, longest_name * 9 + 40))
        column_config["作付名"] = st.column_config.TextColumn(
            "作付名",
            width=crop_name_width,
        )

    if "作付名" in display_df.columns:
        display_df["作付名"] = display_df["作付名"].mask(
            display_df["作付名"].duplicated(),
            "",
        )

    styled_df = style_web_table(display_df, original_df, max_count)

    st.dataframe(
        styled_df,
        use_container_width=True,
        hide_index=True,
        height=900,
        column_config=column_config,
    )


def load_csv_from_same_folder(filename: str):
    """app.py と同じフォルダにCSVがあれば読み込む"""
    local_path = Path(filename)

    if local_path.exists():
        return read_csv_auto(local_path), str(local_path)

    return None, None


# =========================
# 画面
# =========================

st.title("農薬散布記録一覧")

drive_file_ids = get_drive_file_ids()
drive_configured = all(drive_file_ids.values())

try:
    if drive_configured:
        with st.spinner("Google Driveから最新データを読み込んでいます..."):
            raw_df = read_google_drive_csv(drive_file_ids["main"])
            pesticide_info_df = read_google_drive_csv(drive_file_ids["pesticide"])
            registration_info_df = read_google_drive_csv(drive_file_ids["registration"])
    else:
        # ローカルでの動作確認用。デプロイ時はSecretsを設定する。
        raw_df, _ = load_csv_from_same_folder("作業記録４ 農薬.csv")
        pesticide_info_df, _ = load_csv_from_same_folder("農薬情報.csv")
        registration_info_df, _ = load_csv_from_same_folder("農薬登録情報.csv")

        if raw_df is None or pesticide_info_df is None or registration_info_df is None:
            st.error(
                "Google Driveの設定がありません。StreamlitのSecretsに"
                "3つのCSVファイルIDを登録してください。"
            )
            st.code(
                '[google_drive]\n'
                'main_csv_file_id = "作業記録CSVのファイルID"\n'
                'pesticide_info_file_id = "農薬情報CSVのファイルID"\n'
                'registration_info_file_id = "農薬登録情報CSVのファイルID"',
                language="toml",
            )
            st.stop()

    if raw_df is not None and pesticide_info_df is not None and registration_info_df is not None:
        (
            base_df,
            unmatched_pesticide_info_count,
            unmatched_registration_count,
        ) = prepare_base_dataframe(raw_df, pesticide_info_df, registration_info_df)

        min_date = base_df["日付"].min().date()
        max_date = base_df["日付"].max().date()

        default_start = pd.Timestamp(
            year=base_df["日付"].max().year,
            month=4,
            day=1,
        ).date()

        if default_start < min_date:
            default_start = min_date

        st.markdown(f"**更新日：{max_date.strftime('%Y/%m/%d')}**")

        st.success("CSVを読み込みました。")

        if unmatched_pesticide_info_count > 0:
            st.warning(
                f"農薬情報CSVに一致しない農薬IDが {unmatched_pesticide_info_count} 行あります。"
                "一致しないものは農薬グループを「未分類」として表示します。"
            )

        if unmatched_registration_count > 0:
            st.warning(
                f"農薬登録情報CSVに一致しない登録番号が {unmatched_registration_count} 行あります。"
                "一致しないものは希釈倍率・使用時期・使用回数を空白で表示します。"
            )

        st.subheader("対象期間の選択")

        col1, col2 = st.columns(2)

        with col1:
            start_date = st.date_input(
                "開始日",
                value=default_start,
                min_value=min_date,
                max_value=max_date,
            )

        with col2:
            end_date = st.date_input(
                "終了日",
                value=max_date,
                min_value=min_date,
                max_value=max_date,
            )

        if start_date > end_date:
            st.error("開始日は終了日より前の日付にしてください。")
            st.stop()

        summary, max_count = prepare_summary(
            base_df,
            start_date=start_date,
            end_date=end_date,
        )

        available_pesticide_groups = sorted_pesticide_groups(
            summary["農薬グループ"].unique()
        )

        st.subheader("農薬グループの表示切り替え")

        selected_pesticide_groups = []

        checkbox_cols = st.columns(
            min(4, max(1, len(available_pesticide_groups)))
        )

        for i, group_name in enumerate(available_pesticide_groups):
            with checkbox_cols[i % len(checkbox_cols)]:
                checked = st.checkbox(
                    group_name,
                    value=group_name not in {"展着剤", "除草剤"},
                    key=f"pesticide_group_{group_name}",
                )

                if checked:
                    selected_pesticide_groups.append(group_name)

        if not selected_pesticide_groups:
            st.error("表示する農薬グループを1つ以上選択してください。")
            st.stop()

        summary_filtered_by_pesticide_group = summary[
            summary["農薬グループ"].isin(selected_pesticide_groups)
        ].copy()

        st.subheader("表示項目の切り替え")

        selected_optional_columns = []

        option_cols = st.columns(4)

        for i, item in enumerate(OPTIONAL_DISPLAY_COLUMNS):
            with option_cols[i % 4]:
                checked = st.checkbox(
                    item["label"],
                    value=item["key"] != "有効成分",
                    key=f"display_column_{item['key']}",
                )

                if checked:
                    selected_optional_columns.append(item["key"])

        st.caption("※ 作付名と散布回数の日付は常に表示されます。")

        st.caption(
            f"対象期間：{pd.Timestamp(start_date).strftime('%Y/%m/%d')} 〜 "
            f"{pd.Timestamp(end_date).strftime('%Y/%m/%d')}"
        )

        st.caption(
            f"作付名×農薬 件数：{len(summary_filtered_by_pesticide_group)}件 / "
            f"最大散布回数：{max_count}回"
        )

        field_groups = sorted([
            clean_text(g)
            for g in summary_filtered_by_pesticide_group["圃場グループ"].dropna().unique()
            if clean_text(g) != ""
        ])

        if not field_groups:
            st.error("圃場グループが見つかりません。")
            st.stop()

        selected_field_group = st.selectbox(
            "圃場グループを選択",
            options=field_groups,
            index=(
                field_groups.index("北海道農場")
                if "北海道農場" in field_groups
                else 0
            ),
        )

        filtered = summary_filtered_by_pesticide_group[
            summary_filtered_by_pesticide_group["圃場グループ"] == selected_field_group
        ].copy()

        count_cols = [f"{i}回目" for i in range(1, max_count + 1)]
        display_cols = (
            get_display_column_keys(selected_optional_columns)
            + count_cols
            + ["_使用可能回数"]
        )

        st.subheader("グループ別一覧")
        display_table_like_excel(filtered[display_cols], max_count=max_count)

        excel_bytes = create_excel(
            summary,
            max_count,
            start_date=start_date,
            end_date=end_date,
            selected_field_group=selected_field_group,
            selected_pesticide_groups=selected_pesticide_groups,
            selected_optional_columns=selected_optional_columns,
        )

        safe_field_group = clean_text(selected_field_group)
        safe_field_group = safe_field_group.replace("/", "_").replace("\\", "_")

        file_name = (
            f"農薬散布記録一覧_"
            f"{pd.Timestamp(start_date).strftime('%Y%m%d')}_"
            f"{pd.Timestamp(end_date).strftime('%Y%m%d')}_"
            f"{safe_field_group}.xlsx"
        )

        st.download_button(
            label="Excelをダウンロード",
            data=excel_bytes,
            file_name=file_name,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

except Exception as e:
    st.error(f"エラーが発生しました: {e}")
